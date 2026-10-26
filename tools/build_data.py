"""Excel -> data/prep.json + data/dishes.json   (one-off conversion, see tools/README.md)

usage:  python3 tools/build_data.py "<folder with the 'Prep Week N - Day - Recipes.xlsx' files>" <output folder>

Each workbook: sheet 1 = prep list, sheet 2 = dry-store pick sheet (ignored here), sheets 3+ = one recipe per sheet (mains and soups; a VEG row never links to a recipe).
WARNING: the JSON in data/ is now the source of truth. Write to a scratch folder and diff before copying over data/.
"""
import openpyxl, glob, json, re, os, sys, collections

if len(sys.argv) != 3:
    sys.exit(__doc__)
SRC, OUT = sys.argv[1], sys.argv[2]

DAYS = {"monday":"mon","tuesday":"tue","wednesday":"wed","thursday":"thu","friday":"fri"}
LOC_ORDER = ["Meat Fridge","Processed Fridge","Whole Veg Fridge","Dairy Fridge","Freezer","Dry Store","Prep Room"]
LOC_MAP = {"meat fridge":"Meat Fridge","vegetable processed fridge":"Processed Fridge",
           "vegetable whole fridge":"Whole Veg Fridge","dairy fridge":"Dairy Fridge",
           "freezer":"Freezer","dry store":"Dry Store","prep room":"Prep Room"}
UNIT_MAP = {"kilogram":"kg","kg":"kg","gram":"g","g":"g","liter":"L","l":"L","milliliter":"mL","ml":"mL",
            "each":"each","piece":"pc","unit":"","serve":"serve","sheets":"sheets"}

def norm_loc(v):
    if not v: return None
    v = re.sub(r"^\d+\s*-\s*","",str(v)).strip().lower()
    return LOC_MAP.get(v)

def fmt_num(x):
    if isinstance(x,(int,float)):
        return str(int(x)) if float(x).is_integer() else ("%g" % x)
    return str(x)

def fmt_amount(amt, unit):
    if amt is None: return ""
    a = fmt_num(amt)
    if not unit: return a
    u = str(unit).strip()
    if "/" in u:
        base, per = u.split("/",1)
        return f"{a} {base.lower()} ({per})"
    return f"{a} {UNIT_MAP.get(u.lower(), u.lower())}".strip()

def slugify(t):
    t = t.lower().replace("&"," and ").replace("+"," and ")
    t = re.sub(r"[^a-z0-9]+","-",t).strip("-")
    return t

def parse_dish(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    title = rows[0][0].strip()
    sub = (rows[1][0] or "") if len(rows)>1 else ""
    ings, notes, group = [], [], None
    if sub.upper().startswith("BIP"):
        notes.append("Bought in product (no recipe)")
    for r in rows[4:]:
        r = (r + [None]*5)[:5]
        a,b,c,d,e = r
        if all(x is None for x in r): continue
        a_s = a.strip() if isinstance(a,str) else a
        if a_s == "Ingredient": continue
        # notes
        if isinstance(a_s,str) and re.match(r"(?i)^\(?recipe notes?\)?:?\s*$", a_s) and d:
            notes.append(str(d).strip()); continue
        if isinstance(a_s,str) and re.match(r"(?i)^recipe notes?:", a_s):
            notes.append(re.sub(r"(?i)^recipe notes?:\s*","",a_s)); continue
        if (a_s is None or a_s=="" ) and isinstance(d,str):
            notes.append(re.sub(r"(?i)^recipe note:\s*","",d.strip())); continue
        if isinstance(a_s,str) and a_s.upper()=="NOTE" and d:
            notes.append(re.sub(r"^NOTE:?\s*","",str(d).strip())); continue
        # sub heading
        if isinstance(a_s,str) and all(x is None for x in (b,c,d,e)):
            group = re.sub(r"[:]\s*$","",a_s); 
            group = re.sub(r"\s*\(used for.*\)$","",group)
            continue
        ings.append({"group":group,"name":a_s,"amount":fmt_amount(b,c),"note":(str(d).strip() if d else ""),
                     "loc":norm_loc(e)})
    return title, ings, notes

dishes_raw = []   # (title, week, day, ings, notes)
prep = []
for f in sorted(glob.glob(os.path.join(SRC,"*.xlsx"))):
    m = re.search(r"Week (\d) - (\w+) - Recipes", os.path.basename(f))
    week, day = int(m.group(1)), DAYS[m.group(2).lower()]
    wb = openpyxl.load_workbook(f, data_only=True)
    ws0 = wb.worksheets[0]
    sheets = {}
    for s in wb.worksheets[2:]:
        t, ings, notes = parse_dish(s)
        sheets[t] = (ings, notes)
        dishes_raw.append((t, week, day, ings, notes))
    section = None
    for r in ws0.iter_rows(min_row=4, values_only=True):
        a,b,c,d,e = (list(r)+[None]*5)[:5]
        if a in (None,"Item"): continue
        if b is None and c is None and d is None and e is None and a.isupper():
            section = a; continue
        prep.append({"week":week,"day":day,"section":section,"item":a.strip(),"_title":a.strip() if (a.strip() in sheets and section in ("MAINS","SOUPS")) else None,
                     "usageNote":(b or "").strip(),"qty":(c or "").strip() if isinstance(c,str) else (fmt_num(c) if c else ""),
                     "notes":(d or "").strip(),"_locs":set()})

# --- dedupe dishes
by_slug = collections.defaultdict(list)
def sig(ings,notes): return json.dumps([[i["group"],i["name"],i["amount"],i["note"],i["loc"]] for i in ings])
for t,w,d,ings,notes in dishes_raw:
    by_slug[slugify(t)].append((t,w,d,ings,notes))
dishes = []; key_to_slug = {}
conflicts = []
for slug, vers in by_slug.items():
    sigs = collections.OrderedDict()
    for v in vers: sigs.setdefault(sig(v[3],v[4]), []).append(v)
    for si,(s,vs) in enumerate(sigs.items()):
        t,w,d,ings,notes = vs[0]
        if len(sigs)>1:
            use = ", ".join(sorted({f"Wk{x[1]} {x[2].title()}" for x in vs}))
            sl = f"{slug}-wk{w}-{d}"; name = f"{t} ({use} prep)"
            conflicts.append((t,[(x[1],x[2]) for x in vs]))
        else:
            sl, name = slug, t
        dishes.append({"dishSlug":sl,"dishName":name,
            "ingredients":[{**({"group":i["group"]} if i["group"] else {}),"ingredient":i["name"],"amountText":i["amount"],
                            **({"notes":i["note"]} if i["note"] else {})} for i in ings],
            **({"notes":notes} if notes else {}),
            "_locs":{i["loc"] for i in ings if i["loc"]}})
        for x in vs: key_to_slug[(x[0],x[1],x[2])] = (sl,{i["loc"] for i in ings if i["loc"]})

# --- finish prep rows
KW = [("meat fridge","Meat Fridge"),("dry store","Dry Store"),("freezer","Freezer"),("dairy fridge","Dairy Fridge")]
out_prep=[]
for p in prep:
    slug=""; locs=set()
    if p["_title"]:
        slug, locs = key_to_slug[(p["_title"],p["week"],p["day"])]
    low = (p["notes"]+" "+p["qty"]).lower()
    for k,v in KW:
        if k in low: locs.add(v)
    locs &= {"Freezer","Dry Store","Meat Fridge","Dairy Fridge"} if not p["_title"] else {"Freezer","Dry Store"}
    if not p["_title"]: pass
    location = " | ".join(sorted(locs, key=LOC_ORDER.index))
    out_prep.append({"week":p["week"],"day":p["day"],"section":p["section"],"item":p["item"],"dishSlug":slug,
                     "location":location,"usageNote":p["usageNote"],"qty":p["qty"],"notes":p["notes"]})
for d in dishes: d.pop("_locs")
dishes.sort(key=lambda d:d["dishName"].lower())
os.makedirs(OUT,exist_ok=True)
json.dump(out_prep, open(os.path.join(OUT,"prep.json"),"w"), indent=2, ensure_ascii=False)
json.dump(dishes, open(os.path.join(OUT,"dishes.json"),"w"), indent=2, ensure_ascii=False)
print("prep rows",len(out_prep),"| dishes",len(dishes),"| dish links",sum(1 for p in out_prep if p["dishSlug"]))
print("CONFLICTS (same name, different recipe):"); [print("  ",c) for c in conflicts]
# unusual amounts
for d in dishes:
    for i in d["ingredients"]:
        if not i["amountText"] or not re.match(r"^[\d.]+( |$)", i["amountText"]): print("ODD AMOUNT:",d["dishName"],"|",i)
# dish refs unmatched
slugs={d["dishSlug"] for d in dishes}
print("missing slugs:",[p["item"] for p in out_prep if p["dishSlug"] and p["dishSlug"] not in slugs])
print("unused dish sheets:", len(slugs)-len({p["dishSlug"] for p in out_prep if p["dishSlug"]}))
